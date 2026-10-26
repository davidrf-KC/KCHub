import os
import json
import threading
import sqlite3
import datetime
import io
import re
import logging
import subprocess
import tempfile
import pandas as pd
import customtkinter as ctk
import tkinter as tk
from tkinter import filedialog, messagebox
from pathlib import Path
import fitz  # PyMuPDF
from PIL import Image, ImageTk

# --- LIBRERÍAS EXCEL ---
import openpyxl
from copy import copy
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment
from openpyxl.drawing.image import Image as pxImage

# --- MÓDULO DE ACTUALIZACIONES ---
import auto_updater

try:
    import docx
except ImportError:
    docx = None

# --- CONFIGURACIÓN DEL LOG ---
ruta_app_data = Path.home() / "Documents" / "KeyCapital AI Hub"
ruta_app_data.mkdir(parents=True, exist_ok=True)

ruta_log = ruta_app_data / "proceso_excel.log"
logging.basicConfig(
    filename=str(ruta_log),
    level=logging.DEBUG,
    format='%(asctime)s - %(levelname)s - %(message)s',
    filemode='w'
)

class KeyCapitalHub(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("Key Capital Licitaciones (Creado por: David Román F)")
        self.geometry("1200x850")
        ctk.set_appearance_mode("System")
        
        # --- VARIABLES ESTADO: EXCEL ---
        self.excel_ruta_matriz = ""
        self.excel_ruta_plantilla = ""

        # --- VARIABLES ESTADO: GENERAL ---
        self.datos_extraidos = {} 
        self.entradas_edicion = {} 
        self.texto_convocatoria = ""
        
        self.pdf_current_folder = ""
        self.pdf_folders_data = {}
        self.pdf_folder_order = []
        
        # --- VARIABLES ESTADO: FOLIADO ---
        self.foliado_pdf_path = ""
        self.foliado_current_page = 0
        self.foliado_total_pages = 0
        self.foliado_doc = None
        self.foliado_custom_font_path = ""
        self.foliado_zoom = 1.0
        self.foliado_drag_data = {"x": 0, "y": 0, "item": False}
        
        # --- VARIABLES ESTADO: DIVIDIR ---
        self.div_pdf_path = ""
        self.div_current_page = 0
        self.div_total_pages = 0
        self.div_doc = None
        self.div_zoom = 1.0
        self.div_cortes = []
        
        # --- VARIABLES ESTADO: OPTIMIZACIÓN APLANADO ---
        self.opt_pdf_paths = [] 
        self.opt_current_page = 0
        self.opt_total_pages = 0
        self.opt_doc = None
        self.opt_zoom = 1.0
        self.opt_drag_data = {"x": 0, "y": 0, "item": False}

        # --- VARIABLES ESTADO: OPTIMIZACIÓN TEXTO (GS) ---
        self.opt_txt_archivo_original = None
        self.opt_txt_pagina_actual = 1
        self.opt_txt_total_paginas = 1
        self.opt_txt_zoom_factor = 1.0
        self.opt_txt_imagen_pil_actual = None
        self.comando_gs = 'gswin64c' if os.name == 'nt' else 'gs'
        self.flags_gs = subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0
        self.dir_temp = tempfile.gettempdir()
        self.opt_txt_temp_page = os.path.join(self.dir_temp, "pagina_temp.pdf")
        self.opt_txt_temp_compressed = os.path.join(self.dir_temp, "comprimido_temp.pdf")
        self.opt_txt_temp_png = os.path.join(self.dir_temp, "preview_temp.png")
        
        self.columnas_completas = [
            'Generar', 'Licitación', 'Ciudad_Estado', 'Fecha', 'Titular_Convocante', 
            'Nombre_Convocante', 'Tipo_Procedimiento', 'Numero_Procedimiento', 
            'Descripcion_Trabajos', 'Nombre_Licitante', 'Representante_Legal', 
            'Domicilio_Calle', 'Domicilio_Numero', 'Domicilio_Colonia', 
            'Domicilio_Ciudad', 'Domicilio_CP', 'Domicilio_Estado', 'Contacto_Telefono', 
            'Contacto_Fax', 'Contacto_Correo', 'Pagina_Web', 'Nacionalidad_Empresa', 
            'RFC_Empresa', 'RFC_Representante', 'Representante_Legal_firma', 
            'Convocatoria_licitación', 'Procedimiento_de_contratación', 
            'Forma_de_procedimiento', 'Razon_social', 'Domicilio_Ciudad_solo', 
            'Entidad_Federativa', 'Numeral_convocatoria', 
            '12_Numeral_convocatoria_subcontrato', '12_Numeral_convocatoria_base'
        ]
        
        meses = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
        hoy = datetime.datetime.now()
        fecha_hoy = f"{hoy.day} de {meses[hoy.month - 1]} de {hoy.year}"
        
        self.datos_empresa = {
            'Ciudad_Estado': 'Ciudad de México',
            'Fecha': fecha_hoy,
            'Nombre_Licitante': 'KEY CAPITAL S.A.P.I. DE C.V.',
            'Representante_Legal': 'Mario Eduardo González Flores',
            'Domicilio_Calle': 'Avenida Nuevo León',
            'Domicilio_Numero': '257, Piso 7',
            'Domicilio_Colonia': 'Escandón I Sección',
            'Domicilio_Ciudad': 'Alcaldía Miguel Hidalgo',
            'Domicilio_CP': '11800',
            'Domicilio_Estado': 'Ciudad de México, México.',
            'Contacto_Telefono': '557592 8496',
            'Contacto_Correo': 'licitaciones@keycapital.com.mx',
            'Pagina_Web': 'www.keycapital.com.mx',
            'Nacionalidad_Empresa': 'MEXICANA',
            'RFC_Empresa': 'KCA150216KC6',
            'RFC_Representante': 'GOFM830509FA2',
            'Representante_Legal_firma': 'ING. MARIO EDUARDO GONZÁLEZ FLORES'
        }
        
        self.inicializar_bd()
        self.crear_estructura_carpetas()
        
        self.grid_columnconfigure(1, weight=1) 
        self.grid_rowconfigure(0, weight=1)
        self.crear_menu_lateral()
        
        self.frame_main = ctk.CTkFrame(self, corner_radius=0, fg_color="transparent")
        self.frame_main.grid(row=0, column=1, sticky="nsew", padx=10, pady=10)
        self.abrir_modulo_pdf()

        # Comprueba actualizaciones en segundo plano 3 segundos tras iniciar
        self.after(3000, auto_updater.iniciar_busqueda_fondo)

    def inicializar_bd(self):
        ruta_db = ruta_app_data / "hub_config.db"
        self.conn = sqlite3.connect(str(ruta_db), check_same_thread=False)
        self.cursor = self.conn.cursor()
        
        self.cursor.execute('''CREATE TABLE IF NOT EXISTS configuraciones (id INTEGER PRIMARY KEY, ip TEXT, puerto TEXT, modelo TEXT)''')
        self.cursor.execute("SELECT ip, puerto, modelo FROM configuraciones WHERE id=1")
        row = self.cursor.fetchone()
        if row:
            self.ip, self.puerto, self.modelo = row
        else:
            self.ip, self.puerto, self.modelo = "127.0.0.1", "11434", "qwen2.5"
            self.cursor.execute("INSERT INTO configuraciones (id, ip, puerto, modelo) VALUES (1, ?, ?, ?)", (self.ip, self.puerto, self.modelo))
            self.conn.commit()
            
        from ollama import Client
        self.cliente_ia = Client(host=f'http://{self.ip}:{self.puerto}')
        
        self.cursor.execute('''CREATE TABLE IF NOT EXISTS carpetas (id INTEGER PRIMARY KEY, bd TEXT, plantillas TEXT, generados TEXT)''')
        self.cursor.execute("SELECT bd, plantillas, generados FROM carpetas WHERE id=1")
        row_rutas = self.cursor.fetchone()
        ruta_docs_base = Path.home() / "Documents"
        if row_rutas:
            self.carpeta_bd = Path(row_rutas[0])
            self.carpeta_plantillas = Path(row_rutas[1])
            self.carpeta_generados = Path(row_rutas[2])
        else:
            self.carpeta_bd = ruta_docs_base / "BD"
            self.carpeta_plantillas = ruta_docs_base / "Plantillas"
            self.carpeta_generados = ruta_docs_base / "Documentos Generados"
            self.cursor.execute("INSERT INTO carpetas (id, bd, plantillas, generados) VALUES (1, ?, ?, ?)", (str(self.carpeta_bd), str(self.carpeta_plantillas), str(self.carpeta_generados)))
            self.conn.commit()
            
        self.cursor.execute('''CREATE TABLE IF NOT EXISTS config_foliado_v2 (
                                id INTEGER PRIMARY KEY, fuente TEXT, tamano TEXT, color TEXT, 
                                digitos TEXT, pag_inicio TEXT, num_inicio TEXT, pos_x REAL, pos_y REAL)''')
        self.cursor.execute("SELECT fuente, tamano, color, digitos, pag_inicio, num_inicio, pos_x, pos_y FROM config_foliado_v2 WHERE id=1")
        row_foliado = self.cursor.fetchone()
        if row_foliado:
            self.fol_fuente, self.fol_tamano, self.fol_color, self.fol_digitos, self.fol_pag_inicio, self.fol_num_inicio, self.foliado_pos_x, self.foliado_pos_y = row_foliado
        else:
            self.fol_fuente, self.fol_tamano, self.fol_color, self.fol_digitos, self.fol_pag_inicio, self.fol_num_inicio, self.foliado_pos_x, self.foliado_pos_y = ("Helvetica", "12", "Negro", "4", "1", "1", 500, 750)
            self.cursor.execute("INSERT INTO config_foliado_v2 (id, fuente, tamano, color, digitos, pag_inicio, num_inicio, pos_x, pos_y) VALUES (1, ?, ?, ?, ?, ?, ?, ?, ?)", 
                                (self.fol_fuente, self.fol_tamano, self.fol_color, self.fol_digitos, self.fol_pag_inicio, self.fol_num_inicio, self.foliado_pos_x, self.foliado_pos_y))
            self.conn.commit()

    def crear_estructura_carpetas(self):
        for carpeta in [self.carpeta_bd, self.carpeta_plantillas, self.carpeta_generados]:
            carpeta.mkdir(parents=True, exist_ok=True)

    def crear_menu_lateral(self):
        self.frame_menu = ctk.CTkFrame(self, width=200, corner_radius=0)
        self.frame_menu.grid(row=0, column=0, sticky="nsew")
        ctk.CTkLabel(self.frame_menu, text="KC Modules", font=("Segoe UI", 20, "bold")).pack(pady=20, padx=10)
        ctk.CTkButton(self.frame_menu, text="📄 Creación de Docs", command=self.abrir_modulo_pdf).pack(pady=10, padx=20)
        ctk.CTkButton(self.frame_menu, text="🖼️ Herramientas PDF", command=self.abrir_modulo_herramientas_pdf).pack(pady=10, padx=20)
        ctk.CTkButton(self.frame_menu, text="📊 Inyección Excel", command=self.abrir_modulo_excel).pack(pady=10, padx=20)
        
        # Botón para buscar actualizaciones manualmente
        ctk.CTkButton(
            self.frame_menu, 
            text="🔄 Buscar Updates", 
            fg_color="#F39C12", 
            hover_color="#D68910", 
            text_color="white",
            command=lambda: auto_updater.buscar_y_aplicar(silencioso=False)
        ).pack(side="bottom", pady=(0, 20), padx=20)

        ctk.CTkButton(self.frame_menu, text="⚙️ Configuración", command=self.abrir_modulo_config).pack(side="bottom", pady=20, padx=20)

    def actualizar_estado(self, texto, cargando=False):
        if hasattr(self, 'lbl_estado_ia') and self.lbl_estado_ia.winfo_exists():
            self.lbl_estado_ia.configure(text=texto)
            if cargando:
                self.progressbar.pack(side="right", padx=20)
                self.progressbar.start()
            else:
                self.progressbar.stop()
                self.progressbar.pack_forget()

    # ==========================================
    # MÓDULO 1: EXTRACCIÓN Y CREACIÓN DE DOCS
    # ==========================================
    def abrir_modulo_pdf(self):
        for widget in self.frame_main.winfo_children(): widget.destroy()
        
        ctk.CTkLabel(self.frame_main, text="Analizador Automático de Convocatorias", font=("Segoe UI", 24, "bold")).pack(pady=(20, 10), anchor="w", padx=20)
        self.tabview_creacion = ctk.CTkTabview(self.frame_main)
        self.tabview_creacion.pack(fill="both", expand=True, padx=20, pady=5)
        
        tab_extraccion = self.tabview_creacion.add("Extracción y Guardado")
        tab_generar = self.tabview_creacion.add("Generar Documentación")
        
        frame_top_extraccion = ctk.CTkFrame(tab_extraccion, fg_color="transparent")
        frame_top_extraccion.pack(fill="x", padx=10, pady=10)
        
        ctk.CTkButton(frame_top_extraccion, text="📁 Cargar PDF de Convocatoria", width=200, command=self.cargar_pdf).pack(side="left", padx=(0, 10))
        
        self.btn_analizar = ctk.CTkButton(frame_top_extraccion, text="🤖 Extraer Datos", width=150, state="disabled", command=self.analizar_pdf)
        self.btn_analizar.pack(side="left")

        self.btn_guardar_bd = ctk.CTkButton(frame_top_extraccion, text="💾 Guardar en Base de Datos", width=180, font=("Segoe UI", 12, "bold"), command=self.guardar_base_datos)
        self.btn_guardar_bd.pack(side="left", padx=(10, 0))
        
        self.lbl_pdf_estado = ctk.CTkLabel(tab_extraccion, text="Ningún archivo cargado", text_color="gray60")
        self.lbl_pdf_estado.pack(anchor="w", padx=15, pady=(0, 5))
        
        self.frame_estado = ctk.CTkFrame(tab_extraccion, fg_color="transparent")
        self.frame_estado.pack(fill="x", padx=15, pady=5)
        self.lbl_estado_ia = ctk.CTkLabel(self.frame_estado, text="Formulario listo para edición o extracción.")
        self.lbl_estado_ia.pack(side="left")
        self.progressbar = ctk.CTkProgressBar(self.frame_estado, mode="indeterminate", width=250)
        self.progressbar.set(0)
        
        self.frame_resultados = ctk.CTkScrollableFrame(tab_extraccion)
        self.frame_resultados.pack(fill="both", expand=True, padx=10, pady=10)
        
        ctk.CTkLabel(tab_generar, text="Generación Automática de Plantillas", font=("Segoe UI", 18, "bold")).pack(pady=(20, 10), anchor="w", padx=20)
        
        frame_sel_plantillas = ctk.CTkFrame(tab_generar, fg_color="transparent")
        frame_sel_plantillas.pack(fill="x", padx=20, pady=5)
        
        ctk.CTkButton(frame_sel_plantillas, text="📁 Seleccionar Carpeta de Plantillas", command=self.seleccionar_carpeta_plantillas_doc).pack(side="left")
        self.lbl_carpeta_plantillas = ctk.CTkLabel(frame_sel_plantillas, text="Usando carpeta por defecto", text_color="gray60")
        self.lbl_carpeta_plantillas.pack(side="left", padx=10)
        
        self.frame_lista_plantillas = ctk.CTkScrollableFrame(tab_generar, height=200)
        self.frame_lista_plantillas.pack(fill="x", padx=20, pady=10)
        
        self.btn_generar_doc_word = ctk.CTkButton(tab_generar, text="📝 Generar Documentación Word", width=250, height=40, font=("Segoe UI", 14, "bold"), command=self.generar_documentacion)
        self.btn_generar_doc_word.pack(anchor="w", padx=20, pady=10)
        
        self.lbl_estado_generacion = ctk.CTkLabel(tab_generar, text="", text_color="#27AE60", font=("Segoe UI", 12, "bold"), justify="left")
        self.lbl_estado_generacion.pack(anchor="w", padx=20)

        self.ruta_plantillas_actual = self.carpeta_plantillas
        self.actualizar_lista_plantillas()
        self.mostrar_formulario_edicion()

    def cargar_pdf(self):
        ruta = filedialog.askopenfilename(filetypes=[("Archivos PDF", "*.pdf")])
        if not ruta: return
        self.lbl_pdf_estado.configure(text=f"Archivo cargado: {os.path.basename(ruta)}")
        self.update()
        texto_extraido = ""
        try:
            doc = fitz.open(ruta)
            total_pags = len(doc)
            texto_extraido += doc[0].get_text() + "\n"
            if total_pags > 1:
                texto_extraido += "\n\n--- FINAL DEL DOCUMENTO (FIRMAS) ---\n\n"
                texto_extraido += doc[total_pags - 1].get_text() + "\n"
            self.texto_convocatoria = texto_extraido.strip()
            if not self.texto_convocatoria:
                self.btn_analizar.configure(state="disabled")
            else:
                self.btn_analizar.configure(state="normal")
            self.actualizar_estado("Documento listo para analizar.", cargando=False)
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo leer el PDF: {e}")

    def analizar_pdf(self):
        self.btn_analizar.configure(state="disabled")
        self.actualizar_estado("Enviando texto a la Inteligencia Artificial...", cargando=True)
        threading.Thread(target=self._procesar_pdf_ia, daemon=True).start()

    def _procesar_pdf_ia(self):
        prompt = f"""
        Extrae la información clave EXACTAMENTE como está escrita. Devuelve ESTRICTAMENTE un JSON:
        {{
            "Licitación": "", "Titular_Convocante": "", "Nombre_Convocante": "",
            "Tipo_Procedimiento": "", "Numero_Procedimiento": "", "Descripcion_Trabajos": "", "Cargo_Titular": ""
        }}
        Texto: {self.texto_convocatoria[:5000]} 
        """
        try:
            self.after(0, lambda: self.actualizar_estado("La IA está leyendo...", cargando=True))
            respuesta = self.cliente_ia.chat(model=self.modelo, format='json', messages=[{'role': 'user', 'content': prompt}])
            datos_raw = json.loads(respuesta['message']['content'])
            self.datos_extraidos = {k: str(v).strip().replace('\n', ' ').upper() if k=="Descripcion_Trabajos" else str(v).strip().replace('\n', ' ') for k, v in datos_raw.items()}
            self.after(0, self.mostrar_formulario_edicion)
            self.after(0, lambda: self.actualizar_estado("✅ Análisis exitoso.", cargando=False))
        except Exception as e:
            self.after(0, lambda: self.mostrar_error_interfaz(str(e)))
        finally:
            self.after(0, lambda: self.btn_analizar.configure(state="normal"))

    def mostrar_formulario_edicion(self):
        for widget in self.frame_resultados.winfo_children(): widget.destroy()
        self.entradas_edicion = {}
        fila_actual = 1
        for clave in self.columnas_completas:
            if clave == 'Generar': continue
            valor_final = self.datos_extraidos.get(clave, self.datos_empresa.get(clave, ""))
            ctk.CTkLabel(self.frame_resultados, text=f"{clave.replace('_', ' ')}:", font=("Segoe UI", 12, "bold")).grid(row=fila_actual, column=0, padx=(10, 5), pady=5, sticky="ne")
            ent = ctk.CTkTextbox(self.frame_resultados, height=75 if clave in ["Descripcion_Trabajos"] else 28, wrap="word" if clave in ["Descripcion_Trabajos"] else "none")
            ent.insert("1.0", str(valor_final))
            ent.grid(row=fila_actual, column=1, padx=(5, 10), pady=5, sticky="ew")
            self.frame_resultados.grid_columnconfigure(1, weight=1)
            self.entradas_edicion[clave] = ent
            fila_actual += 1

    def mostrar_error_interfaz(self, error_str):
        for widget in self.frame_resultados.winfo_children(): widget.destroy()
        ctk.CTkLabel(self.frame_resultados, text=f"❌ Error:\n{error_str}", text_color="#E74C3C").grid(row=0, column=0, sticky="w")
        self.actualizar_estado("❌ Error de red.", cargando=False)

    def guardar_base_datos(self):
        if not self.entradas_edicion: return
        datos = {k: ("" if k=='Generar' else v.get("1.0", "end-1c").strip().replace('\n',' ')) for k, v in self.entradas_edicion.items()}
        ruta = filedialog.asksaveasfilename(initialdir=self.carpeta_bd, initialfile=f"{datos.get('Licitación', 'Licitacion')}.xlsx", defaultextension=".xlsx")
        if ruta:
            pd.DataFrame([datos]).to_excel(ruta, index=False)
            messagebox.showinfo("Éxito", "Matriz guardada.")

    def seleccionar_carpeta_plantillas_doc(self):
        carpeta = filedialog.askdirectory(initialdir=self.ruta_plantillas_actual)
        if carpeta:
            self.ruta_plantillas_actual = Path(carpeta)
            self.lbl_carpeta_plantillas.configure(text=f"Carpeta: {self.ruta_plantillas_actual.name}")
            self.actualizar_lista_plantillas()

    def actualizar_lista_plantillas(self):
        for widget in self.frame_lista_plantillas.winfo_children():
            widget.destroy()
            
        if not hasattr(self, 'ruta_plantillas_actual') or not self.ruta_plantillas_actual.exists():
            ctk.CTkLabel(self.frame_lista_plantillas, text="La ruta especificada no existe.", text_color="red").pack(anchor="w")
            self.btn_generar_doc_word.configure(state="disabled")
            return

        plantillas = [p for p in self.ruta_plantillas_actual.rglob("*.docx") if not p.name.startswith("~$")]
        if not plantillas:
            ctk.CTkLabel(self.frame_lista_plantillas, text="No se encontraron archivos .docx en esta carpeta.", text_color="#E74C3C").pack(anchor="w", padx=5)
            self.btn_generar_doc_word.configure(state="disabled")
        else:
            self.btn_generar_doc_word.configure(state="normal")
            ctk.CTkLabel(self.frame_lista_plantillas, text=f"Se detectaron {len(plantillas)} archivos listos para procesar:", font=("Segoe UI", 12, "bold"), text_color="#2980B9").pack(anchor="w", padx=5, pady=(0,5))
            for p in plantillas:
                rel_path = p.relative_to(self.ruta_plantillas_actual)
                ctk.CTkLabel(self.frame_lista_plantillas, text=f"📄 {rel_path}").pack(anchor="w", padx=15)

    def generar_documentacion(self):
        if docx is None: 
            messagebox.showerror("Librería Faltante", "No está instalada la librería python-docx.\nAsegúrate de tenerla instalada para generar los documentos.")
            return
            
        ruta_excel = filedialog.askopenfilename(title="Selecciona la Matriz (Excel) a inyectar", initialdir=self.carpeta_bd, filetypes=[("Excel", "*.xlsx")])
        if not ruta_excel: return
        
        self.lbl_estado_generacion.configure(text="Generando documentos, por favor espera...", text_color="#F39C12")
        self.update()
        
        try:
            datos = pd.read_excel(ruta_excel).iloc[0].fillna("").to_dict()
            plantillas = [p for p in self.ruta_plantillas_actual.rglob("*.docx") if not p.name.startswith("~$")]
            
            nombre_proyecto = str(datos.get('Licitación', 'Licitacion_Generada')).replace('/','_').replace('\\','_')
            carpeta_destino = self.carpeta_generados / nombre_proyecto
            
            docs_generados = 0
            for ruta in plantillas:
                destino = carpeta_destino / ruta.relative_to(self.ruta_plantillas_actual)
                destino.parent.mkdir(parents=True, exist_ok=True)
                doc = docx.Document(ruta)
                for p in doc.paragraphs: self._reemplazar_texto(p, datos)
                for t in doc.tables:
                    for f in t.rows:
                        for c in f.cells:
                            for p in c.paragraphs: self._reemplazar_texto(p, datos)
                doc.save(destino)
                docs_generados += 1
                
            msg_exito = f"¡Éxito! Se generaron {docs_generados} documentos.\nGuardados en: {carpeta_destino}"
            self.lbl_estado_generacion.configure(text=msg_exito, text_color="#27AE60")
            messagebox.showinfo("Proceso Completado", f"Documentación generada exitosamente.\n\nRuta de acceso:\n{carpeta_destino}")
            
        except Exception as e:
            self.lbl_estado_generacion.configure(text="Error durante la generación.", text_color="#E74C3C")
            messagebox.showerror("Error", f"Ocurrió un error al procesar las plantillas:\n{str(e)}")

    def _reemplazar_texto(self, parrafo, mapa):
        if not parrafo.text: return
        for k, v in mapa.items():
            for et in [f"«{k}»", f"<<{k}>>", f"{{{{{k}}}}}"]:
                if et in parrafo.text:
                    r = False
                    for run in parrafo.runs:
                        if et in run.text:
                            run.text = run.text.replace(et, str(v))
                            r = True
                    if not r or et in parrafo.text: parrafo.text = parrafo.text.replace(et, str(v))

    # ==========================================
    # MÓDULO 2: HERRAMIENTAS PDF
    # ==========================================
    def abrir_modulo_herramientas_pdf(self):
        for widget in self.frame_main.winfo_children(): widget.destroy()
        
        ctk.CTkLabel(self.frame_main, text="Herramientas PDF", font=("Segoe UI", 24, "bold")).pack(pady=(20, 10), anchor="w", padx=20)
        self.tabview_pdf = ctk.CTkTabview(self.frame_main)
        self.tabview_pdf.pack(fill="both", expand=True, padx=20, pady=5)
        
        tab_word = self.tabview_pdf.add("Word a PDF")
        tab_fusion = self.tabview_pdf.add("Fusión")
        tab_foliar = self.tabview_pdf.add("Foliar")
        tab_dividir = self.tabview_pdf.add("Dividir")
        tab_optimizar = self.tabview_pdf.add("Optimización por Aplanado")
        tab_opt_texto = self.tabview_pdf.add("Optimización por Texto (GS)")
        
        ctk.CTkButton(tab_word, text="1. Cargar Carpeta", command=self.pdf_seleccionar_carpeta_word).pack(pady=10)
        self.lbl_word_estado = ctk.CTkLabel(tab_word, text="Ninguna carpeta")
        self.lbl_word_estado.pack()
        self.btn_convertir_word = ctk.CTkButton(tab_word, text="2. Iniciar Conversión", state="disabled", command=self.pdf_convertir_word_a_pdf)
        self.btn_convertir_word.pack(pady=10)

        ctk.CTkButton(tab_fusion, text="Seleccionar Carpeta", command=self.pdf_select_folder).pack(pady=10)
        self.lbl_pdf_path = ctk.CTkLabel(tab_fusion, text="Ninguna")
        self.lbl_pdf_path.pack()
        ctk.CTkButton(tab_fusion, text="Unir PDFs", command=self.pdf_process_pdfs).pack(pady=10)
        self.pdf_scroll_frame = ctk.CTkScrollableFrame(tab_fusion)
        self.pdf_scroll_frame.pack(fill="both", expand=True)

        main_foliado = ctk.CTkFrame(tab_foliar, fg_color="transparent")
        main_foliado.pack(fill="both", expand=True)
        left_foliado = ctk.CTkFrame(main_foliado, width=380)
        left_foliado.pack(side="left", fill="y", padx=10, pady=10)
        
        ctk.CTkButton(left_foliado, text="Cargar PDF a Foliar", command=self.foliado_cargar_pdf).pack(pady=15, padx=15, fill="x")
        self.lbl_foliado_archivo = ctk.CTkLabel(left_foliado, text="Ningún archivo", text_color="gray60")
        self.lbl_foliado_archivo.pack(pady=(0,10))
        
        frame_opts = ctk.CTkFrame(left_foliado, fg_color="transparent")
        frame_opts.pack(fill="x", padx=15)
        
        ctk.CTkLabel(frame_opts, text="Fuente:", font=("Segoe UI", 12, "bold")).grid(row=0, column=0, sticky="w", pady=5)
        frame_f = ctk.CTkFrame(frame_opts, fg_color="transparent")
        frame_f.grid(row=0, column=1, sticky="w", pady=5)
        self.fol_combo_fuente = ctk.CTkComboBox(frame_f, values=["Helvetica", "Times-Roman", "Courier", "Personalizada"], width=130, command=self.foliado_on_font_change)
        self.fol_combo_fuente.set(self.fol_fuente)
        self.fol_combo_fuente.pack(side="left", padx=(0,5))
        self.btn_cargar_fuente = ctk.CTkButton(frame_f, text="Cargar .ttf", width=70, command=self.foliado_cargar_fuente)
        self.btn_cargar_fuente.pack(side="left")
        if self.fol_fuente != "Personalizada": 
            self.btn_cargar_fuente.configure(state="disabled")
        
        self.lbl_fuente_cargada = ctk.CTkLabel(frame_opts, text="", text_color="#27AE60", font=("Segoe UI", 10))
        self.lbl_fuente_cargada.grid(row=1, column=1, sticky="w")
        
        ctk.CTkLabel(frame_opts, text="Tamaño (px):", font=("Segoe UI", 12, "bold")).grid(row=2, column=0, sticky="w", pady=5)
        self.fol_entry_tamano = ctk.CTkEntry(frame_opts, width=80)
        self.fol_entry_tamano.insert(0, self.fol_tamano)
        self.fol_entry_tamano.grid(row=2, column=1, sticky="w", pady=5)
        
        ctk.CTkLabel(frame_opts, text="Color:", font=("Segoe UI", 12, "bold")).grid(row=3, column=0, sticky="w", pady=5)
        self.fol_combo_color = ctk.CTkComboBox(frame_opts, values=["Negro", "Gris", "Rojo", "Azul"], width=100)
        self.fol_combo_color.set(self.fol_color)
        self.fol_combo_color.grid(row=3, column=1, sticky="w", pady=5)
        
        ctk.CTkLabel(frame_opts, text="Ceros de relleno:", font=("Segoe UI", 12, "bold")).grid(row=4, column=0, sticky="w", pady=5)
        self.fol_combo_digitos = ctk.CTkComboBox(frame_opts, values=["1", "2", "3", "4", "5", "6"], width=80)
        self.fol_combo_digitos.set(self.fol_digitos)
        self.fol_combo_digitos.grid(row=4, column=1, sticky="w", pady=5)
        
        ctk.CTkLabel(frame_opts, text="Foliar desde pág.:", font=("Segoe UI", 12, "bold")).grid(row=5, column=0, sticky="w", pady=5)
        self.fol_entry_pag_inicio = ctk.CTkEntry(frame_opts, width=80)
        self.fol_entry_pag_inicio.insert(0, self.fol_pag_inicio)
        self.fol_entry_pag_inicio.grid(row=5, column=1, sticky="w", pady=5)
        
        ctk.CTkLabel(frame_opts, text="Empezar folio en:", font=("Segoe UI", 12, "bold")).grid(row=6, column=0, sticky="w", pady=5)
        self.fol_entry_num_inicio = ctk.CTkEntry(frame_opts, width=80)
        self.fol_entry_num_inicio.insert(0, self.fol_num_inicio)
        self.fol_entry_num_inicio.grid(row=6, column=1, sticky="w", pady=5)

        ctk.CTkLabel(left_foliado, text="* Arrastra el número en la vista previa a la derecha para colocarlo donde desees.", text_color="#F39C12", font=("Segoe UI", 11, "italic"), wraplength=280).pack(pady=(15,10))
        
        ctk.CTkButton(left_foliado, text="Actualizar Aspecto", command=self.foliado_actualizar_preview).pack(pady=5)
        self.btn_ejecutar_foliado = ctk.CTkButton(left_foliado, text="Generar PDF Foliado", state="disabled", command=self.foliado_procesar)
        self.btn_ejecutar_foliado.pack(pady=10)
        
        right_foliado = ctk.CTkFrame(main_foliado)
        right_foliado.pack(side="right", fill="both", expand=True)
        top_bar_preview = ctk.CTkFrame(right_foliado, fg_color="transparent")
        top_bar_preview.pack(fill="x", pady=10)
        ctk.CTkButton(top_bar_preview, text="◀", width=40, command=lambda: self.foliado_cambiar_pagina(-1)).pack(side="left")
        self.lbl_foliado_paginacion = ctk.CTkLabel(top_bar_preview, text="Pág. 0 de 0")
        self.lbl_foliado_paginacion.pack(side="left", expand=True)
        ctk.CTkButton(top_bar_preview, text="▶", width=40, command=lambda: self.foliado_cambiar_pagina(1)).pack(side="right")
        self.canvas_preview = tk.Canvas(right_foliado, bg="#2b2b2b")
        self.canvas_preview.pack(fill="both", expand=True)
        
        self.canvas_preview.tag_bind("movable", "<ButtonPress-1>", self.on_drag_start)
        self.canvas_preview.tag_bind("movable", "<B1-Motion>", self.on_drag_motion)

        main_div = ctk.CTkFrame(tab_dividir, fg_color="transparent")
        main_div.pack(fill="both", expand=True)
        left_div = ctk.CTkFrame(main_div, width=350)
        left_div.pack(side="left", fill="y", padx=10, pady=10)
        ctk.CTkButton(left_div, text="Cargar PDF a Dividir", command=self.div_cargar_pdf).pack(pady=15, padx=15, fill="x")
        self.lbl_div_archivo = ctk.CTkLabel(left_div, text="Ningún archivo")
        self.lbl_div_archivo.pack(pady=(0,10))
        ctk.CTkButton(left_div, text="✂️ Agregar corte aquí", command=self.div_agregar_corte).pack(pady=10)
        self.frame_div_cortes = ctk.CTkScrollableFrame(left_div)
        self.frame_div_cortes.pack(fill="both", expand=True, pady=10)
        self.btn_ejecutar_div = ctk.CTkButton(left_div, text="Generar Segmentados", state="disabled", command=self.div_procesar)
        self.btn_ejecutar_div.pack(pady=15)
        
        right_div = ctk.CTkFrame(main_div)
        right_div.pack(side="right", fill="both", expand=True)
        top_bar_div = ctk.CTkFrame(right_div, fg_color="transparent")
        top_bar_div.pack(fill="x", pady=10)
        ctk.CTkButton(top_bar_div, text="◀", width=40, command=lambda: self.div_cambiar_pagina(-1)).pack(side="left")
        self.ent_div_pagina = ctk.CTkEntry(top_bar_div, width=40)
        self.ent_div_pagina.pack(side="left", padx=10)
        self.ent_div_pagina.bind("<Return>", self.div_ir_a_pagina)
        self.lbl_div_total_paginas = ctk.CTkLabel(top_bar_div, text=" de 0")
        self.lbl_div_total_paginas.pack(side="left")
        ctk.CTkButton(top_bar_div, text="▶", width=40, command=lambda: self.div_cambiar_pagina(1)).pack(side="right")
        self.canvas_div_preview = tk.Canvas(right_div, bg="#2b2b2b")
        self.canvas_div_preview.pack(fill="both", expand=True)

        main_opt = ctk.CTkFrame(tab_optimizar, fg_color="transparent")
        main_opt.pack(fill="both", expand=True)
        left_opt = ctk.CTkFrame(main_opt, width=350)
        left_opt.pack(side="left", fill="y", padx=10, pady=10)
        ctk.CTkLabel(left_opt, text="Optimización a JPG", font=("Segoe UI", 16, "bold")).pack(anchor="w", pady=10, padx=15)
        ctk.CTkButton(left_opt, text="📁 Cargar PDFs", command=self.opt_cargar_pdf).pack(fill="x", padx=15, pady=5)
        self.lbl_opt_archivo = ctk.CTkLabel(left_opt, text="Ningún archivo", text_color="gray50")
        self.lbl_opt_archivo.pack(anchor="w", padx=15, pady=10)
        
        ctk.CTkLabel(left_opt, text="Resolución (DPI):", font=("Segoe UI", 12, "bold")).pack(anchor="w", padx=15)
        self.opt_dpi_var = ctk.StringVar(value="150")
        ctk.CTkComboBox(left_opt, variable=self.opt_dpi_var, values=["72", "100", "150", "200"], command=lambda _: self.opt_actualizar_preview()).pack(fill="x", padx=15, pady=5)
        
        ctk.CTkLabel(left_opt, text="Calidad (%):", font=("Segoe UI", 12, "bold")).pack(anchor="w", padx=15)
        self.opt_calidad_var = ctk.StringVar(value="75")
        ctk.CTkComboBox(left_opt, variable=self.opt_calidad_var, values=["40", "50", "65", "75", "85"], command=lambda _: self.opt_actualizar_preview()).pack(fill="x", padx=15, pady=5)
        
        self.btn_opt_procesar = ctk.CTkButton(left_opt, text="⚡ Iniciar Compresión", state="disabled", command=self.opt_procesar_pdf)
        self.btn_opt_procesar.pack(fill="x", padx=15, pady=20)
        self.lbl_opt_estado = ctk.CTkLabel(left_opt, text="", text_color="#27AE60")
        self.lbl_opt_estado.pack(anchor="w", padx=15, pady=10)

        right_opt = ctk.CTkFrame(main_opt, fg_color=("gray85", "gray18"))
        right_opt.pack(side="right", fill="both", expand=True, padx=(0,10), pady=10)
        top_bar_opt = ctk.CTkFrame(right_opt, fg_color="transparent")
        top_bar_opt.pack(fill="x", pady=10, padx=10)
        ctk.CTkButton(top_bar_opt, text="◀ Anterior", width=80, command=lambda: self.opt_cambiar_pagina(-1)).pack(side="left")
        self.lbl_opt_paginacion = ctk.CTkLabel(top_bar_opt, text="Pág. 0 de 0", font=("Segoe UI", 12, "bold"))
        self.lbl_opt_paginacion.pack(side="left", expand=True)
        ctk.CTkButton(top_bar_opt, text="Siguiente ▶", width=80, command=lambda: self.opt_cambiar_pagina(1)).pack(side="right")
        self.frame_opt_canvas = ctk.CTkFrame(right_opt)
        self.frame_opt_canvas.pack(fill="both", expand=True, padx=10, pady=10)
        self.canvas_opt_preview = tk.Canvas(self.frame_opt_canvas, bg="#2b2b2b")
        self.canvas_opt_preview.pack(fill="both", expand=True)
        
        self.canvas_opt_preview.bind("<ButtonPress-1>", self.opt_on_drag_start)
        self.canvas_opt_preview.bind("<B1-Motion>", self.opt_on_drag_motion)

        main_opt_txt = ctk.CTkFrame(tab_opt_texto, fg_color="transparent")
        main_opt_txt.pack(fill="both", expand=True)
        
        left_opt_txt = ctk.CTkFrame(main_opt_txt, width=350)
        left_opt_txt.pack(side="left", fill="y", padx=10, pady=10)
        
        ctk.CTkLabel(left_opt_txt, text="Compresor Avanzado (GS)", font=("Segoe UI", 16, "bold"), wraplength=300).pack(anchor="w", pady=(10,5), padx=15)
        ctk.CTkLabel(left_opt_txt, text="Ideal para contratos pesados. Conserva el texto real y comprime imágenes de fondo.", text_color="gray60", wraplength=300, justify="left").pack(anchor="w", pady=(0, 20), padx=15)
        
        ctk.CTkButton(left_opt_txt, text="1. Seleccionar PDF", command=self.opt_txt_seleccionar_archivo).pack(fill="x", padx=15, pady=5)
        self.lbl_opt_txt_archivo = ctk.CTkLabel(left_opt_txt, text="Ningún archivo seleccionado", text_color="gray50", wraplength=300)
        self.lbl_opt_txt_archivo.pack(anchor="w", padx=15, pady=(0, 15))
        
        ctk.CTkLabel(left_opt_txt, text="2. Resolución Máxima (DPI):", font=("Segoe UI", 12, "bold")).pack(anchor="w", padx=15, pady=(5,0))
        self.opt_txt_slider_calidad = ctk.CTkSlider(left_opt_txt, from_=50, to=300, number_of_steps=25, command=self.opt_txt_actualizar_label_slider)
        self.opt_txt_slider_calidad.set(120)
        self.opt_txt_slider_calidad.configure(state="disabled")
        self.opt_txt_slider_calidad.pack(fill="x", padx=15, pady=5)
        self.lbl_opt_txt_slider_val = ctk.CTkLabel(left_opt_txt, text="120 DPI (Recomendado)")
        self.lbl_opt_txt_slider_val.pack(anchor="w", padx=15, pady=(0,15))
        
        self.btn_opt_txt_actualizar = ctk.CTkButton(left_opt_txt, text="Actualizar Vista Previa", state="disabled", command=self.opt_txt_procesar_vista_previa)
        self.btn_opt_txt_actualizar.pack(fill="x", padx=15, pady=5)
        
        self.lbl_opt_txt_estado = ctk.CTkLabel(left_opt_txt, text="", font=("Segoe UI", 12, "bold"))
        self.lbl_opt_txt_estado.pack(anchor="w", padx=15, pady=10)
        
        self.btn_opt_txt_guardar = ctk.CTkButton(left_opt_txt, text="3. Guardar PDF Comprimido", state="disabled", command=self.opt_txt_guardar_pdf_completo)
        self.btn_opt_txt_guardar.pack(fill="x", padx=15, pady=20)

        right_opt_txt = ctk.CTkFrame(main_opt_txt, fg_color=("gray85", "gray18"))
        right_opt_txt.pack(side="right", fill="both", expand=True, padx=(0,10), pady=10)
        
        top_bar_opt_txt = ctk.CTkFrame(right_opt_txt, fg_color="transparent")
        top_bar_opt_txt.pack(fill="x", pady=10, padx=10)
        
        self.btn_opt_txt_prev = ctk.CTkButton(top_bar_opt_txt, text="◀", width=40, state="disabled", command=self.opt_txt_pagina_anterior)
        self.btn_opt_txt_prev.pack(side="left")
        self.lbl_opt_txt_pagina = ctk.CTkLabel(top_bar_opt_txt, text="Página - de -", font=("Segoe UI", 12, "bold"))
        self.lbl_opt_txt_pagina.pack(side="left", expand=True)
        self.btn_opt_txt_next = ctk.CTkButton(top_bar_opt_txt, text="▶", width=40, state="disabled", command=self.opt_txt_pagina_siguiente)
        self.btn_opt_txt_next.pack(side="left")
        
        ctk.CTkButton(top_bar_opt_txt, text="Zoom -", width=60, command=self.opt_txt_zoom_out).pack(side="right", padx=5)
        self.lbl_opt_txt_zoom = ctk.CTkLabel(top_bar_opt_txt, text="100%", width=40)
        self.lbl_opt_txt_zoom.pack(side="right")
        ctk.CTkButton(top_bar_opt_txt, text="Zoom +", width=60, command=self.opt_txt_zoom_in).pack(side="right", padx=5)
        
        self.frame_opt_txt_canvas = ctk.CTkFrame(right_opt_txt)
        self.frame_opt_txt_canvas.pack(fill="both", expand=True, padx=10, pady=10)
        
        self.canvas_opt_txt = tk.Canvas(self.frame_opt_txt_canvas, bg="#2b2b2b", highlightthickness=0)
        hbar_txt = tk.Scrollbar(self.frame_opt_txt_canvas, orient="horizontal", command=self.canvas_opt_txt.xview)
        vbar_txt = tk.Scrollbar(self.frame_opt_txt_canvas, orient="vertical", command=self.canvas_opt_txt.yview)
        self.canvas_opt_txt.configure(xscrollcommand=hbar_txt.set, yscrollcommand=vbar_txt.set)
        vbar_txt.pack(side="right", fill="y")
        hbar_txt.pack(side="bottom", fill="x")
        self.canvas_opt_txt.pack(side="left", fill="both", expand=True)
        self.canvas_opt_txt_img_id = self.canvas_opt_txt.create_image(0, 0, anchor="nw")

    def pdf_seleccionar_carpeta_word(self):
        folder = filedialog.askdirectory()
        if folder:
            self.ruta_word_origen = folder
            self.lbl_word_estado.configure(text=f"Carpeta: {folder}")
            self.btn_convertir_word.configure(state="normal")
            
    def pdf_convertir_word_a_pdf(self):
        threading.Thread(target=self._hilo_convertir_word, daemon=True).start()

    def _hilo_convertir_word(self):
        import pythoncom, comtypes.client
        pythoncom.CoInitialize() 
        try:
            origen = Path(self.ruta_word_origen)
            destino = origen.parent / f"{origen.name}-PDF"
            archivos = [(Path(root)/f, destino/(Path(root)/f).relative_to(origen).with_suffix('.pdf')) for root, _, files in os.walk(origen) for f in files if f.endswith(('.docx', '.doc')) and not f.startswith('~$')]
            word = comtypes.client.CreateObject('Word.Application')
            word.Visible = False
            for src, dst in archivos:
                dst.parent.mkdir(parents=True, exist_ok=True)
                doc = word.Documents.Open(str(src))
                doc.SaveAs(str(dst), FileFormat=17) 
                doc.Close()
            word.Quit()
            self.after(0, lambda: messagebox.showinfo("Terminado", "Conversión finalizada exitosamente."))
        except Exception: pass
        finally: pythoncom.CoUninitialize()

    def pdf_select_folder(self):
        folder = filedialog.askdirectory()
        if folder:
            self.pdf_current_folder = folder
            self.lbl_pdf_path.configure(text=f"Carpeta: {folder}")
            self.pdf_folders_data = {}
            self.pdf_folder_order = []
            for widget in self.pdf_scroll_frame.winfo_children(): widget.destroy()
            archivos_encontrados = 0
            for root, _, files in os.walk(folder):
                pdf_files = [os.path.join(root, f) for f in sorted(files) if f.lower().endswith(".pdf")]
                if pdf_files: 
                    rel_path = os.path.relpath(root, folder)
                    dn = f"📁 {os.path.basename(folder)} (Carpeta Principal)" if rel_path == "." else f"📁 {rel_path}"
                    self.pdf_folders_data[dn] = {"abs_path": root, "files": pdf_files}
                    self.pdf_folder_order.append(dn)
                    ctk.CTkLabel(self.pdf_scroll_frame, text=dn, font=("Segoe UI", 14, "bold"), text_color="#2980B9").pack(anchor="w", pady=(10, 0), padx=5)
                    for f in pdf_files:
                        ctk.CTkLabel(self.pdf_scroll_frame, text=f"   📄 {os.path.basename(f)}").pack(anchor="w", padx=10)
                        archivos_encontrados += 1
            if archivos_encontrados == 0:
                ctk.CTkLabel(self.pdf_scroll_frame, text="No se encontraron archivos PDF en esta carpeta.", text_color="#E74C3C").pack(pady=10)

    def pdf_process_pdfs(self):
        if not self.pdf_folder_order: return
        origen_root = Path(self.pdf_current_folder)
        destino_root = origen_root.parent / f"{origen_root.name}-Unidos"
        try:
            for folder_name in self.pdf_folder_order:
                fi = self.pdf_folders_data[folder_name]
                if not fi["files"]: continue
                sub = Path(fi["abs_path"])
                target_dir = destino_root / sub.relative_to(origen_root)
                target_dir.mkdir(parents=True, exist_ok=True)
                output = os.path.join(target_dir, f"{sub.name if sub.name else 'Documento'}.pdf")
                doc_final = fitz.open()
                for path in fi["files"]:
                    d = fitz.open(path)
                    doc_final.insert_pdf(d)
                    d.close()
                doc_final.save(output)
                doc_final.close()
            messagebox.showinfo("¡Proceso finalizado!", f"Se unieron los PDFs correctamente.\nGuardados en:\n{destino_root}")
        except Exception as e:
            messagebox.showerror("Error", f"Ocurrió un error al unir los PDFs:\n{str(e)}")

    def foliado_cargar_pdf(self):
        ruta = filedialog.askopenfilename(filetypes=[("PDF", "*.pdf")])
        if ruta:
            self.foliado_pdf_path = ruta
            self.lbl_foliado_archivo.configure(text=Path(ruta).name)
            if self.foliado_doc: self.foliado_doc.close()
            self.foliado_doc = fitz.open(ruta)
            self.foliado_total_pages = len(self.foliado_doc)
            self.btn_ejecutar_foliado.configure(state="normal")
            self.foliado_actualizar_preview()

    def foliado_on_font_change(self, val):
        if val == "Personalizada":
            self.btn_cargar_fuente.configure(state="normal")
        else:
            self.btn_cargar_fuente.configure(state="disabled")
            self.lbl_fuente_cargada.configure(text="")
            self.foliado_custom_font_path = ""

    def foliado_cargar_fuente(self):
        ruta = filedialog.askopenfilename(filetypes=[("Fuentes TrueType", "*.ttf")])
        if ruta:
            self.foliado_custom_font_path = ruta
            self.lbl_fuente_cargada.configure(text=Path(ruta).name)
            self.foliado_actualizar_preview()

    def _get_foliado_color_rgb(self, nombre):
        if nombre == "Rojo": return (1, 0, 0)
        elif nombre == "Azul": return (0, 0, 1)
        elif nombre == "Gris": return (0.5, 0.5, 0.5)
        return (0, 0, 0)

    def _get_foliado_color_hex(self, nombre):
        if nombre == "Rojo": return "#FF0000"
        elif nombre == "Azul": return "#0000FF"
        elif nombre == "Gris": return "#808080"
        return "#000000"

    def foliado_cambiar_pagina(self, delta):
        if self.foliado_doc and 0 <= self.foliado_current_page + delta < self.foliado_total_pages:
            self.foliado_current_page += delta
            self.foliado_actualizar_preview()

    def on_drag_start(self, event):
        self.foliado_drag_data["x"] = self.canvas_preview.canvasx(event.x)
        self.foliado_drag_data["y"] = self.canvas_preview.canvasy(event.y)
        self.foliado_drag_data["item"] = True

    def on_drag_motion(self, event):
        if self.foliado_drag_data["item"]:
            cx, cy = self.canvas_preview.canvasx(event.x), self.canvas_preview.canvasy(event.y)
            self.canvas_preview.move("movable", cx - self.foliado_drag_data["x"], cy - self.foliado_drag_data["y"])
            self.foliado_drag_data["x"], self.foliado_drag_data["y"] = cx, cy
            coords = self.canvas_preview.coords("folio_text")
            if coords:
                self.foliado_pos_x, self.foliado_pos_y = coords[0] / self.foliado_zoom, coords[1] / self.foliado_zoom

    def foliado_actualizar_preview(self):
        if not self.foliado_doc: return
        try:
            self.canvas_preview.delete("all")
            page = self.foliado_doc[self.foliado_current_page]
            matrix = fitz.Matrix(self.foliado_zoom, self.foliado_zoom)
            pix = page.get_pixmap(matrix=matrix)
            mode = "RGBA" if pix.alpha else "RGB"
            img = Image.frombytes(mode, [pix.width, pix.height], pix.samples)
            self.tk_img_preview = ImageTk.PhotoImage(img) 
            self.canvas_preview.create_image(0, 0, image=self.tk_img_preview, anchor="nw")
            self.canvas_preview.configure(scrollregion=(0, 0, pix.width, pix.height))
            
            try: digitos = int(self.fol_combo_digitos.get())
            except: digitos = 4
            try: pag_inicio_pdf = int(self.fol_entry_pag_inicio.get())
            except: pag_inicio_pdf = 1
            try: num_inicial = int(self.fol_entry_num_inicio.get())
            except: num_inicial = 1
            
            pagina_actual_fisica = self.foliado_current_page + 1
            if pagina_actual_fisica >= pag_inicio_pdf:
                numero_actual = num_inicial + (pagina_actual_fisica - pag_inicio_pdf)
                texto_folio = str(numero_actual).zfill(digitos)
                fontsize = int(self.fol_entry_tamano.get())
                color_hex = self._get_foliado_color_hex(self.fol_combo_color.get())
                cx = self.foliado_pos_x * self.foliado_zoom
                cy = self.foliado_pos_y * self.foliado_zoom
                fs_canvas = int(fontsize * self.foliado_zoom)
                if fs_canvas < 6: fs_canvas = 6
                
                f_combo = self.fol_combo_fuente.get()
                tk_font = "Arial"
                if f_combo == "Times-Roman": tk_font = "Times"
                elif f_combo == "Courier": tk_font = "Courier"

                self.canvas_preview.create_text(
                    cx, cy,
                    text=texto_folio,
                    fill=color_hex,
                    font=(tk_font, fs_canvas, "bold"),
                    anchor="sw", 
                    tags=("movable", "folio_text")
                )
                
                bbox = self.canvas_preview.bbox("folio_text")
                if bbox:
                    pad = 5
                    self.canvas_preview.create_rectangle(
                        bbox[0]-pad, bbox[1]-pad, bbox[2]+pad, bbox[3]+pad,
                        outline="red", dash=(4,4), tags=("movable", "folio_box")
                    )

            self.lbl_foliado_paginacion.configure(text=f"Pág. {pagina_actual_fisica} de {self.foliado_total_pages}")
        except Exception as e:
            messagebox.showerror("Error en Previsualización", str(e))

    def foliado_procesar(self):
        ruta_salida = Path(self.foliado_pdf_path).parent / f"{Path(self.foliado_pdf_path).stem}-Foliado.pdf"
        doc = fitz.open(self.foliado_pdf_path)
        f_combo = self.fol_combo_fuente.get()
        c_rgb = self._get_foliado_color_rgb(self.fol_combo_color.get())
        s_px = int(self.fol_entry_tamano.get())
        
        fontname_int = "helv"
        if f_combo == "Times-Roman": fontname_int = "ti-ro"
        elif f_combo == "Courier": fontname_int = "cour"
        
        for i, page in enumerate(doc):
            if i + 1 >= int(self.fol_entry_pag_inicio.get()):
                txt = str(int(self.fol_entry_num_inicio.get()) + (i + 1 - int(self.fol_entry_pag_inicio.get()))).zfill(int(self.fol_combo_digitos.get()))
                if f_combo == "Personalizada" and self.foliado_custom_font_path:
                    page.insert_font(fontname="F0", fontfile=self.foliado_custom_font_path)
                    page.insert_text(fitz.Point(self.foliado_pos_x, self.foliado_pos_y), txt, fontname="F0", fontsize=s_px, color=c_rgb)
                else:
                    page.insert_text(fitz.Point(self.foliado_pos_x, self.foliado_pos_y), txt, fontname=fontname_int, fontsize=s_px, color=c_rgb)
                    
        doc.save(str(ruta_salida))
        messagebox.showinfo("Completado", f"Archivo generado exitosamente en:\n{ruta_salida}")

    def div_cargar_pdf(self):
        ruta = filedialog.askopenfilename(filetypes=[("PDF", "*.pdf")])
        if ruta:
            self.div_pdf_path = ruta
            self.lbl_div_archivo.configure(text=Path(ruta).name)
            if self.div_doc: self.div_doc.close()
            self.div_doc = fitz.open(ruta)
            self.div_total_pages = len(self.div_doc)
            self.div_current_page = 0
            self.div_cortes = [1]
            self.btn_ejecutar_div.configure(state="normal")
            self.div_actualizar_preview()
            self.div_renderizar_cortes()
            
    def div_cambiar_pagina(self, delta):
        if self.div_doc and 0 <= self.div_current_page + delta < self.div_total_pages:
            self.div_current_page += delta
            self.div_actualizar_preview()
            
    def div_ir_a_pagina(self, event):
        try:
            if 1 <= int(self.ent_div_pagina.get()) <= self.div_total_pages:
                self.div_current_page = int(self.ent_div_pagina.get()) - 1
                self.div_actualizar_preview()
        except: pass

    def div_agregar_corte(self):
        if self.div_current_page + 1 not in self.div_cortes:
            self.div_cortes.append(self.div_current_page + 1)
            self.div_cortes.sort()
            self.div_actualizar_preview()
            self.div_renderizar_cortes()

    def div_eliminar_corte(self, pagina):
        if pagina in self.div_cortes:
            self.div_cortes.remove(pagina)
            self.div_actualizar_preview()
            self.div_renderizar_cortes()

    def div_renderizar_cortes(self):
        for widget in self.frame_div_cortes.winfo_children():
            widget.destroy()
        if not self.div_doc: return

        ctk.CTkLabel(self.frame_div_cortes, text="Segmentos a generar:", font=("Segoe UI", 12, "bold"), text_color="#2980B9").pack(anchor="w", pady=(0, 5))
        for i in range(len(self.div_cortes)):
            inicio = self.div_cortes[i]
            fin = self.div_cortes[i+1] - 1 if i + 1 < len(self.div_cortes) else self.div_total_pages
            frame_seg = ctk.CTkFrame(self.frame_div_cortes, fg_color="transparent")
            frame_seg.pack(fill="x", pady=2)
            texto = f"📄 Doc {i+1}: Pág. {inicio} a {fin}"
            ctk.CTkLabel(frame_seg, text=texto).pack(side="left")
            if inicio != 1:
                btn_del = ctk.CTkButton(frame_seg, text="❌", width=25, height=25, fg_color="transparent", hover_color="#E74C3C", text_color="#E74C3C", command=lambda p=inicio: self.div_eliminar_corte(p))
                btn_del.pack(side="right")

    def div_actualizar_preview(self):
        if not self.div_doc: return
        self.canvas_div_preview.delete("all")
        pix = self.div_doc[self.div_current_page].get_pixmap()
        img = Image.frombytes("RGBA" if pix.alpha else "RGB", [pix.width, pix.height], pix.samples)
        self.tk_img_div_preview = ImageTk.PhotoImage(img)
        self.canvas_div_preview.create_image(0, 0, image=self.tk_img_div_preview, anchor="nw")
        self.ent_div_pagina.delete(0, "end")
        self.ent_div_pagina.insert(0, str(self.div_current_page + 1))
        self.lbl_div_total_paginas.configure(text=f" de {self.div_total_pages}")
        if self.div_current_page + 1 in self.div_cortes and self.div_current_page + 1 != 1:
            self.canvas_div_preview.create_line(0, 5, pix.width, 5, fill="red", width=3)

    def div_procesar(self):
        orig = Path(self.div_pdf_path)
        dest = orig.parent / f"{orig.stem}-segmentado"
        dest.mkdir(exist_ok=True)
        doc = fitz.open(self.div_pdf_path)
        cortes = sorted(list(set(self.div_cortes)))
        for idx, start in enumerate(cortes):
            end = cortes[idx+1]-1 if idx+1 < len(cortes) else len(doc)
            doc_n = fitz.open()
            doc_n.insert_pdf(doc, from_page=start-1, to_page=end-1)
            doc_n.save(str(dest / f"{str(idx+1).zfill(2)}-{orig.stem}.pdf"))
        messagebox.showinfo("Completado", f"Generados en {dest}")

    def opt_cargar_pdf(self):
        rutas = filedialog.askopenfilenames(filetypes=[("Archivos PDF", "*.pdf")])
        if rutas:
            try:
                self.opt_pdf_paths = list(rutas)
                self.opt_doc = fitz.open(self.opt_pdf_paths[0])
                self.opt_total_pages = len(self.opt_doc)
                self.opt_current_page = 0
                self.opt_zoom = 1.0
                self.btn_opt_procesar.configure(state="normal")
                self.lbl_opt_archivo.configure(text=f"{len(self.opt_pdf_paths)} archivos seleccionados")
                self.opt_actualizar_preview()
            except Exception: pass

    def opt_cambiar_pagina(self, delta):
        if self.opt_doc and 0 <= self.opt_current_page + delta < self.opt_total_pages:
            self.opt_current_page += delta
            self.opt_actualizar_preview()

    def opt_cambiar_zoom(self, delta):
        if not self.opt_doc: return
        self.opt_zoom += delta
        if self.opt_zoom < 0.2: self.opt_zoom = 0.2
        if self.opt_zoom > 3.0: self.opt_zoom = 3.0
        self.opt_actualizar_preview()

    def opt_on_drag_start(self, event):
        self.opt_drag_data["x"] = self.canvas_opt_preview.canvasx(event.x)
        self.opt_drag_data["y"] = self.canvas_opt_preview.canvasy(event.y)
        self.opt_drag_data["item"] = True

    def opt_on_drag_motion(self, event):
        if self.opt_drag_data["item"]:
            cx = self.canvas_opt_preview.canvasx(event.x)
            cy = self.canvas_opt_preview.canvasy(event.y)
            dx = cx - self.opt_drag_data["x"]
            dy = cy - self.opt_drag_data["y"]
            self.canvas_opt_preview.xview_scroll(int(-dx/10), "units")
            self.canvas_opt_preview.yview_scroll(int(-dy/10), "units")

    def opt_actualizar_preview(self):
        if not self.opt_doc: return
        try:
            dpi_val = int(self.opt_dpi_var.get())
            page = self.opt_doc[self.opt_current_page]
            pix = page.get_pixmap(dpi=dpi_val)
            mode = "RGBA" if pix.alpha else "RGB"
            img = Image.frombytes(mode, [pix.width, pix.height], pix.samples)
            if img.mode != "RGB": img = img.convert("RGB")
            img_byte_arr = io.BytesIO()
            img.save(img_byte_arr, format='JPEG', quality=int(self.opt_calidad_var.get()))
            img_byte_arr.seek(0)
            img_compressed = Image.open(img_byte_arr)
            if self.opt_zoom != 1.0:
                new_w = int(img_compressed.width * self.opt_zoom)
                new_h = int(img_compressed.height * self.opt_zoom)
                img_compressed = img_compressed.resize((new_w, new_h), Image.Resampling.LANCZOS)
            self.tk_img_opt_preview = ImageTk.PhotoImage(img_compressed)
            self.canvas_opt_preview.delete("all")
            self.canvas_opt_preview.create_image(0, 0, image=self.tk_img_opt_preview, anchor="nw")
            self.lbl_opt_paginacion.configure(text=f"Pág. {self.opt_current_page + 1} de {self.opt_total_pages}")
        except Exception: pass

    def opt_procesar_pdf(self):
        if getattr(self, 'opt_pdf_paths', None):
            self.btn_opt_procesar.configure(state="disabled", text="Optimizando...")
            self.lbl_opt_estado.configure(text="Convirtiendo páginas en segundo plano...")
            threading.Thread(target=self._hilo_optimizar, daemon=True).start()

    def _hilo_optimizar(self):
        dpi_val = int(self.opt_dpi_var.get())
        calidad_jpg = int(self.opt_calidad_var.get())
        total_ahorro_mb = 0
        archivos_procesados = 0
        try:
            for idx, ruta in enumerate(self.opt_pdf_paths):
                origen = Path(ruta)
                destino = origen.parent / f"{origen.stem}-Optimizado.pdf"
                self.after(0, lambda r=origen.name, i=idx+1, t=len(self.opt_pdf_paths): self.lbl_opt_estado.configure(text=f"Optimizando ({i}/{t}): {r}..."))
                tamano_inicial = os.path.getsize(origen)
                doc_in = fitz.open(origen)
                doc_out = fitz.open()
                for i in range(len(doc_in)):
                    page = doc_in[i]
                    pix = page.get_pixmap(dpi=dpi_val)
                    mode = "RGBA" if pix.alpha else "RGB"
                    img = Image.frombytes(mode, [pix.width, pix.height], pix.samples)
                    if img.mode != "RGB": img = img.convert("RGB")
                    img_byte_arr = io.BytesIO()
                    img.save(img_byte_arr, format='JPEG', quality=calidad_jpg)
                    new_page = doc_out.new_page(width=page.rect.width, height=page.rect.height)
                    new_page.insert_image(new_page.rect, stream=img_byte_arr.getvalue())
                doc_out.save(str(destino), garbage=4, deflate=True)
                doc_in.close()
                doc_out.close()
                ahorro_mb = (tamano_inicial - os.path.getsize(destino)) / (1024 * 1024)
                total_ahorro_mb += ahorro_mb
                archivos_procesados += 1
            self.after(0, lambda: self.lbl_opt_estado.configure(text=f"¡Exitosa! Reducción: {total_ahorro_mb:.2f} MB.", text_color="#27AE60"))
            self.after(0, lambda: messagebox.showinfo("Completado", f"Se optimizaron {archivos_procesados} archivo(s).\nEspacio ahorrado: {total_ahorro_mb:.2f} MB"))
        except Exception as e:
            self.after(0, lambda: messagebox.showerror("Error", str(e)))
            self.after(0, lambda: self.lbl_opt_estado.configure(text="Error.", text_color="red"))
        finally:
            self.after(0, lambda: self.btn_opt_procesar.configure(state="normal", text="⚡ Iniciar Compresión"))

    def opt_txt_actualizar_label_slider(self, value):
        self.lbl_opt_txt_slider_val.configure(text=f"{int(value)} DPI")

    def opt_txt_seleccionar_archivo(self):
        ruta = filedialog.askopenfilename(title="Selecciona el PDF", filetypes=[("Archivos PDF", "*.pdf")])
        if ruta:
            self.opt_txt_archivo_original = ruta
            self.lbl_opt_txt_archivo.configure(text=f"Archivo: {os.path.basename(ruta)}")
            self.opt_txt_slider_calidad.configure(state="normal")
            self.btn_opt_txt_actualizar.configure(state="normal")
            self.btn_opt_txt_guardar.configure(state="normal")
            try:
                doc_temp = fitz.open(ruta)
                self.opt_txt_total_paginas = len(doc_temp)
                doc_temp.close()
            except:
                self.opt_txt_total_paginas = 1
            self.opt_txt_pagina_actual = 1
            self.opt_txt_zoom_factor = 1.0
            self.opt_txt_actualizar_controles_navegacion()
            self.opt_txt_procesar_vista_previa()

    def opt_txt_actualizar_controles_navegacion(self):
        self.lbl_opt_txt_pagina.configure(text=f"Pág. {self.opt_txt_pagina_actual} de {self.opt_txt_total_paginas}")
        self.btn_opt_txt_prev.configure(state="normal" if self.opt_txt_pagina_actual > 1 else "disabled")
        self.btn_opt_txt_next.configure(state="normal" if self.opt_txt_pagina_actual < self.opt_txt_total_paginas else "disabled")
        self.lbl_opt_txt_zoom.configure(text=f"{int(self.opt_txt_zoom_factor * 100)}%")

    def opt_txt_pagina_anterior(self):
        if self.opt_txt_pagina_actual > 1:
            self.opt_txt_pagina_actual -= 1
            self.opt_txt_actualizar_controles_navegacion()
            self.opt_txt_procesar_vista_previa()

    def opt_txt_pagina_siguiente(self):
        if self.opt_txt_pagina_actual < self.opt_txt_total_paginas:
            self.opt_txt_pagina_actual += 1
            self.opt_txt_actualizar_controles_navegacion()
            self.opt_txt_procesar_vista_previa()

    def opt_txt_zoom_in(self):
        if self.opt_txt_imagen_pil_actual and self.opt_txt_zoom_factor < 3.0:
            self.opt_txt_zoom_factor += 0.25
            self.opt_txt_actualizar_controles_navegacion()
            self.opt_txt_renderizar_imagen()

    def opt_txt_zoom_out(self):
        if self.opt_txt_imagen_pil_actual and self.opt_txt_zoom_factor > 0.25:
            self.opt_txt_zoom_factor -= 0.25
            self.opt_txt_actualizar_controles_navegacion()
            self.opt_txt_renderizar_imagen()

    def opt_txt_obtener_comando_optimizacion(self, archivo_entrada, archivo_salida, dpi):
        return [
            self.comando_gs, 
            "-sDEVICE=pdfwrite", 
            "-dCompatibilityLevel=1.4",
            "-dNOPAUSE", 
            "-dBATCH", 
            "-dQUIET",
            "-dDownsampleColorImages=true", f"-dColorImageResolution={dpi}", "-dColorImageDownsampleType=/Bicubic",
            "-dDownsampleGrayImages=true", f"-dGrayImageResolution={dpi}", "-dGrayImageDownsampleType=/Bicubic",
            "-dDownsampleMonoImages=true", f"-dMonoImageResolution={dpi}", "-dMonoImageDownsampleType=/Bicubic",
            "-dAutoFilterColorImages=false", "-dColorImageFilter=/DCTEncode", 
            "-dAutoFilterGrayImages=false", "-dGrayImageFilter=/DCTEncode",
            f"-sOutputFile={archivo_salida}", 
            archivo_entrada
        ]

    def opt_txt_procesar_vista_previa(self):
        if not self.opt_txt_archivo_original: return
        self.btn_opt_txt_actualizar.configure(state="disabled")
        self.lbl_opt_txt_estado.configure(text="Procesando vista previa (GS)...", text_color="#2980B9")
        threading.Thread(target=self._hilo_opt_txt_preview, daemon=True).start()

    def _hilo_opt_txt_preview(self):
        try:
            dpi = int(self.opt_txt_slider_calidad.get())
            comando_extraer = [
                self.comando_gs, "-sDEVICE=pdfwrite", "-dNOPAUSE", "-dBATCH", "-dQUIET",
                f"-dFirstPage={self.opt_txt_pagina_actual}", f"-dLastPage={self.opt_txt_pagina_actual}",
                f"-sOutputFile={self.opt_txt_temp_page}", self.opt_txt_archivo_original
            ]
            subprocess.run(comando_extraer, check=True, creationflags=self.flags_gs)

            comando_comprimir = self.opt_txt_obtener_comando_optimizacion(self.opt_txt_temp_page, self.opt_txt_temp_compressed, dpi)
            subprocess.run(comando_comprimir, check=True, creationflags=self.flags_gs)

            comando_png = [
                self.comando_gs, "-sDEVICE=png16m", "-r150", "-dNOPAUSE", "-dBATCH", "-dQUIET",
                f"-sOutputFile={self.opt_txt_temp_png}", self.opt_txt_temp_compressed
            ]
            subprocess.run(comando_png, check=True, creationflags=self.flags_gs)

            self.opt_txt_imagen_pil_actual = Image.open(self.opt_txt_temp_png)
            peso_orig = os.path.getsize(self.opt_txt_temp_page)
            peso_comp = os.path.getsize(self.opt_txt_temp_compressed)
            if peso_orig > 0:
                ahorro = ((peso_orig - peso_comp) / peso_orig) * 100
                color = "#27AE60" if ahorro > 0 else "#E67E22"
                self.after(0, lambda a=ahorro, c=color: self.lbl_opt_txt_estado.configure(text=f"Ahorro estimado pág. actual: {a:.1f}%", text_color=c))

            self.after(0, self.opt_txt_renderizar_imagen)
        except FileNotFoundError:
            self.after(0, lambda: self.lbl_opt_txt_estado.configure(text="Error: Ghostscript no encontrado.", text_color="red"))
            self.after(0, lambda: messagebox.showwarning("Dependencia faltante", "No se encontró Ghostscript (gswin64c) en el sistema. Asegúrate de instalarlo para usar esta función."))
        except Exception as e:
            self.after(0, lambda: self.lbl_opt_txt_estado.configure(text="Error al generar vista previa.", text_color="red"))
        finally:
            self.after(0, lambda: self.btn_opt_txt_actualizar.configure(state="normal"))

    def opt_txt_renderizar_imagen(self):
        if not self.opt_txt_imagen_pil_actual: return
        nuevo_ancho = int(self.opt_txt_imagen_pil_actual.width * self.opt_txt_zoom_factor)
        nuevo_alto = int(self.opt_txt_imagen_pil_actual.height * self.opt_txt_zoom_factor)
        img_redimensionada = self.opt_txt_imagen_pil_actual.resize((nuevo_ancho, nuevo_alto), Image.Resampling.LANCZOS)
        self.opt_txt_img_tk = ImageTk.PhotoImage(img_redimensionada)
        self.canvas_opt_txt.itemconfig(self.canvas_opt_txt_img_id, image=self.opt_txt_img_tk)
        self.canvas_opt_txt.config(scrollregion=self.canvas_opt_txt.bbox(tk.ALL))

    def opt_txt_guardar_pdf_completo(self):
        if not self.opt_txt_archivo_original: return
        directorio, nombre_archivo = os.path.split(self.opt_txt_archivo_original)
        nombre, extension = os.path.splitext(nombre_archivo)
        ruta_salida = filedialog.asksaveasfilename(
            title="Guardar PDF Comprimido",
            initialdir=directorio,
            initialfile=f"{nombre}-comprimido{extension}",
            defaultextension=".pdf",
            filetypes=[("Archivos PDF", "*.pdf")]
        )
        if not ruta_salida: return
        self.btn_opt_txt_guardar.configure(state="disabled", text="Comprimiendo...")
        self.lbl_opt_txt_estado.configure(text="Procesando documento completo, espera...", text_color="#2980B9")
        threading.Thread(target=self._hilo_opt_txt_guardar, args=(ruta_salida,), daemon=True).start()

    def _hilo_opt_txt_guardar(self, ruta_salida):
        try:
            dpi = int(self.opt_txt_slider_calidad.get())
            comando = self.opt_txt_obtener_comando_optimizacion(self.opt_txt_archivo_original, ruta_salida, dpi)
            subprocess.run(comando, check=True, creationflags=self.flags_gs)
            peso_original = os.path.getsize(self.opt_txt_archivo_original) / (1024 * 1024)
            peso_final = os.path.getsize(ruta_salida) / (1024 * 1024)
            reduccion = ((peso_original - peso_final) / peso_original) * 100
            msg = (f"¡Compresión exitosa!\n\n"
                   f"Peso original: {peso_original:.2f} MB\n"
                   f"Peso final: {peso_final:.2f} MB\n"
                   f"Reducción total: {reduccion:.2f}%")
            self.after(0, lambda: messagebox.showinfo("Proceso Terminado", msg))
            self.after(0, lambda: self.lbl_opt_txt_estado.configure(text=f"¡Listo! Reducción lograda: {reduccion:.1f}%", text_color="#27AE60"))
        except subprocess.CalledProcessError as e:
            self.after(0, lambda: messagebox.showerror("Error", f"Ocurrió un error al procesar el archivo completo.\n{e}"))
            self.after(0, lambda: self.lbl_opt_txt_estado.configure(text="Error en compresión.", text_color="red"))
        finally:
            self.after(0, lambda: self.btn_opt_txt_guardar.configure(state="normal", text="3. Guardar PDF Comprimido"))

    # ==========================================
    # MÓDULO 4: INYECCIÓN EXCEL
    # ==========================================
    def abrir_modulo_excel(self):
        for widget in self.frame_main.winfo_children(): widget.destroy()
        ctk.CTkLabel(self.frame_main, text="Inyección Dinámica en Excel", font=("Segoe UI", 24, "bold")).pack(pady=(20, 20), anchor="w", padx=20)
        frame_excel = ctk.CTkFrame(self.frame_main, fg_color="transparent")
        frame_excel.pack(fill="both", expand=True, padx=20, pady=10)

        self.excel_rutas_plantillas = []
        self.btn_excel_matriz = ctk.CTkButton(frame_excel, text="1. Cargar Matriz de Datos", command=self.excel_cargar_matriz)
        self.btn_excel_matriz.pack(pady=10, anchor="w")
        self.lbl_excel_matriz = ctk.CTkLabel(frame_excel, text="Ningún archivo seleccionado", text_color="gray60")
        self.lbl_excel_matriz.pack(pady=(0, 20), anchor="w")

        self.btn_excel_plantilla = ctk.CTkButton(frame_excel, text="2. Cargar Plantillas Base", command=self.excel_cargar_plantilla)
        self.btn_excel_plantilla.pack(pady=10, anchor="w")
        self.lbl_excel_plantilla = ctk.CTkLabel(frame_excel, text="Ningún archivo seleccionado", text_color="gray60")
        self.lbl_excel_plantilla.pack(pady=(0, 20), anchor="w")

        self.btn_excel_procesar = ctk.CTkButton(frame_excel, text="3. Generar Reportes", state="disabled", command=self.excel_procesar)
        self.btn_excel_procesar.pack(pady=30, anchor="w")

    def excel_cargar_matriz(self):
        ruta = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx *.xls")])
        if ruta:
            self.excel_ruta_matriz = ruta
            self.lbl_excel_matriz.configure(text=os.path.basename(ruta))
            self.excel_validar_ejecucion()

    def excel_cargar_plantilla(self):
        rutas = filedialog.askopenfilenames(filetypes=[("Excel", "*.xlsx")])
        if rutas:
            self.excel_rutas_plantillas = list(rutas)
            if len(self.excel_rutas_plantillas) == 1:
                self.lbl_excel_plantilla.configure(text=os.path.basename(self.excel_rutas_plantillas[0]))
            else:
                self.lbl_excel_plantilla.configure(text=f"{len(self.excel_rutas_plantillas)} archivos seleccionados")
            self.excel_validar_ejecucion()

    def excel_validar_ejecucion(self):
        if self.excel_ruta_matriz and hasattr(self, 'excel_rutas_plantillas') and len(self.excel_rutas_plantillas) > 0:
            self.btn_excel_procesar.configure(state="normal")

    def excel_reparar_bordes_fusion(self, ws, min_row, min_col, max_row, max_col):
        top_left_border = copy(ws.cell(row=min_row, column=min_col).border)
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                ws.cell(row=row, column=col).border = top_left_border

    def excel_procesar(self):
        self.btn_excel_procesar.configure(state="disabled", text="Procesando...")
        self.update()
        threading.Thread(target=self._hilo_excel_procesar, daemon=True).start()

    def _hilo_excel_procesar(self):
        try:
            logging.info("--- INICIO DE PROCESO MULTIBLOQUE (LOTE) ---")
            wb_matriz = openpyxl.load_workbook(self.excel_ruta_matriz, rich_text=True, data_only=True)
            ws_matriz = wb_matriz.active
            
            datos_bd = []
            for row in range(2, ws_matriz.max_row + 1):
                registro = {}
                is_empty = True
                for col in range(1, ws_matriz.max_column + 1):
                    letra = get_column_letter(col)
                    c_m = ws_matriz.cell(row=row, column=col)
                    val = c_m.value
                    registro[letra] = {
                        "value": val if val is not None else "",
                        "number_format": c_m.number_format,
                        "font": copy(c_m.font) if c_m.font else None
                    }
                    if val is not None and str(val).strip() != "":
                        is_empty = False
                if not is_empty:
                    datos_bd.append(registro)
            wb_matriz.close()
            
            total_archivos = len(self.excel_rutas_plantillas)
            imagenes_cacheadas = {}
            
            for index_archivo, ruta_plantilla in enumerate(self.excel_rutas_plantillas):
                self.after(0, lambda i=index_archivo+1, t=total_archivos: self.btn_excel_procesar.configure(text=f"Procesando {i}/{t}..."))
                wb = openpyxl.load_workbook(ruta_plantilla)
                ws = wb.active

                bloques_config = []
                for row in range(1, ws.max_row + 1):
                    tiene_start = False
                    f_col, f_key = None, None
                    for col in range(1, ws.max_column + 1):
                        val = ws.cell(row=row, column=col).value
                        if isinstance(val, str) and '<<start>>' in val.lower():
                            tiene_start = True
                            tokens = re.findall(r'<<([^>]+)>>', val)
                            for t in tokens:
                                if ',' in t and 'sep=' not in t.lower():
                                    partes = t.split(',', 1)
                                    f_col, f_key = partes[0].strip().upper(), partes[1].strip()
                            break
                    
                    if tiene_start:
                        mapa_instrucciones = {}
                        span_datos = 1
                        filas_separacion = 0
                        for col in range(1, ws.max_column + 1):
                            val = ws.cell(row=row, column=col).value
                            if isinstance(val, str) and ('<<' in val or '(' in val):
                                tokens_brackets = re.findall(r'<<([^>]+)>>', val)
                                parentesis_span = re.findall(r'\(\d+\)', val)
                                col_span = int(parentesis_span[0]) if parentesis_span else 1
                                tags_data = []

                                for t in tokens_brackets:
                                    t_clean, t_lower = t.strip(), t.strip().lower()
                                    if 'sep=' in t_lower:
                                        try: filas_separacion = int(t_lower.split('sep=')[1])
                                        except ValueError: pass
                                    elif t_lower in ['start', 'end'] or ',' in t_clean: continue
                                    elif t_clean.isdigit(): col_span = int(t_clean)
                                    else: tags_data.append(t_clean.upper())
                                        
                                if col_span > span_datos: span_datos = col_span
                                texto_base = re.sub(r'<<start>>|<<end>>', '', str(val), flags=re.IGNORECASE)
                                texto_base = re.sub(r'\(\d+\)', '', texto_base)
                                texto_base = re.sub(r'<<[^>]+,[^>]+>>', '', texto_base)
                                mapa_instrucciones[col] = {'tags': tags_data, 'span': col_span, 'texto_base': texto_base}

                        bloques_config.append({
                            'fila_base': row,
                            'mapa': mapa_instrucciones,
                            'master_span': max(span_datos, 1) + filas_separacion,
                            'max_col_bloque': ws.max_column,
                            'filter_col': f_col, 'filter_keyword': f_key
                        })

                filas_tabla = {r for cfg in bloques_config for r in range(cfg['fila_base'], cfg['fila_base'] + cfg['master_span'])}
                if datos_bd:
                    primer_registro = datos_bd[0]
                    for row_idx in range(1, ws.max_row + 1):
                        if row_idx in filas_tabla: continue
                        for col_idx in range(1, ws.max_column + 1):
                            celda = ws.cell(row=row_idx, column=col_idx)
                            val = celda.value
                            if isinstance(val, str) and '<<' in val and '>>' in val:
                                nuevo_val = val
                                for et in re.findall(r'<<([^>]+)>>', val):
                                    et_upper = et.strip().upper()
                                    if et_upper in primer_registro:
                                        dc = primer_registro[et_upper]
                                        v_crudo = dc.get("value", "")
                                        if isinstance(v_crudo, (datetime.datetime, datetime.date)): tx_rep = v_crudo.strftime('%d/%m/%Y')
                                        elif isinstance(v_crudo, (int, float)):
                                            n_fmt = str(dc.get("number_format", ""))
                                            if '$' in n_fmt or '\"$\"' in n_fmt: tx_rep = f"${v_crudo:,.2f}"
                                            elif '0.00' in n_fmt: tx_rep = f"{v_crudo:,.2f}"
                                            else: tx_rep = str(v_crudo)
                                        else: tx_rep = str(v_crudo)
                                        
                                        if val.strip() == f"<<{et}>>":
                                            nuevo_val = v_crudo
                                            if dc.get("number_format"): celda.number_format = dc.get("number_format")
                                            if dc.get("font"): celda.font = dc.get("font")
                                        else: nuevo_val = nuevo_val.replace(f"<<{et}>>", tx_rep)
                                celda.value = nuevo_val

                desp = 0
                for config in bloques_config:
                    f_base = config['fila_base'] + desp
                    mapa = config['mapa']
                    m_span = config['master_span']
                    m_col = config['max_col_bloque']
                    
                    bd_filtra = []
                    for r in datos_bd:
                        if config['filter_col'] and config['filter_keyword']:
                            if config['filter_keyword'].lower() in str(r.get(config['filter_col'], {}).get("value", "")).lower():
                                bd_filtra.append(r)
                        else: bd_filtra.append(r)

                    for off in range(m_span): ws.row_dimensions[f_base + off].height = None 

                    if len(bd_filtra) > 1:
                        t_filas = (len(bd_filtra) - 1) * m_span
                        ws.insert_rows(f_base + m_span, amount=t_filas)
                        for idx in range(1, len(bd_filtra)):
                            for off in range(m_span):
                                o_row, n_row = f_base + off, f_base + (idx * m_span) + off
                                ws.row_dimensions[n_row].height = None 
                                for col in range(1, m_col + 1):
                                    sc, dc = ws.cell(row=o_row, column=col), ws.cell(row=n_row, column=col)
                                    if sc.has_style:
                                        dc.font, dc.fill, dc.number_format, dc.alignment, dc.border = copy(sc.font), copy(sc.fill), copy(sc.number_format), copy(sc.alignment), copy(sc.border)
                                    v = sc.value
                                    dc.value = re.sub(r'<<[^>]+>>|\(\d+\)', '', v).strip() if (isinstance(v, str) and '<<' in v) else v
                        desp += t_filas

                    for col in range(1, m_col + 1):
                        for off in range(m_span):
                            c = ws.cell(row=f_base + off, column=col)
                            if isinstance(c.value, str) and '<<' in c.value:
                                c.value = re.sub(r'<<[^>]+>>|\(\d+\)', '', c.value).strip() or None

                    for i, reg in enumerate(bd_filtra):
                        f_act = f_base + (i * m_span)
                        max_ln = 1
                        for c_num, info in mapa.items():
                            if info['tags']:
                                txt_val = "".join([str(reg.get(t, {}).get("value", "")) + "\n" for t in info['tags']])
                                a_col = ws.column_dimensions[get_column_letter(c_num)].width or 13.0
                                c_linea = max(int(a_col * 0.9), 5)
                                lns = sum((len(p)//c_linea) + (1 if len(p)%c_linea>0 else 0) if len(p)>0 else 1 for p in txt_val.split('\n'))
                                ln_dist = lns / max(info.get('span', 1), 1)
                                if ln_dist > max_ln: max_ln = ln_dist
                        
                        for off in range(m_span): ws.row_dimensions[f_act + off].height = max_ln * 15.0 + 5.0 
                        
                        for c_num, info in mapa.items():
                            if info['tags']:
                                cdest = ws.cell(row=f_act, column=c_num)
                                if len(info['tags']) == 1 and info['texto_base'].strip() == f"<<{info['tags'][0]}>>":
                                    tag = info['tags'][0]
                                    cdest.value = reg.get(tag, {}).get("value", "")
                                    if reg.get(tag, {}).get("number_format"): cdest.number_format = reg.get(tag, {})["number_format"]
                                    if reg.get(tag, {}).get("font"): cdest.font = reg.get(tag, {})["font"]
                                else:
                                    t_final = info['texto_base']
                                    for tag in info['tags']:
                                        vc = reg.get(tag, {}).get("value", "")
                                        if isinstance(vc, (datetime.datetime, datetime.date)): vr = vc.strftime('%d/%m/%Y')
                                        elif isinstance(vc, (int, float)):
                                            nf = str(reg.get(tag, {}).get("number_format", ""))
                                            vr = f"${vc:,.2f}" if '$' in nf else (f"{vc:,.2f}" if '0.00' in nf else str(vc))
                                        else: vr = str(vc)
                                        t_final = t_final.replace(f"<<{tag}>>", vr)
                                    cdest.value = t_final
                                    if reg.get(info['tags'][0], {}).get("font"): cdest.font = reg.get(info['tags'][0], {})["font"]
                                
                                ca = cdest.alignment
                                cdest.alignment = Alignment(horizontal=ca.horizontal or 'center', vertical='center', wrap_text=True)
                        
                        for c_num, info in mapa.items():
                            t_span = max(info['span'], 1) if info.get('span') else 1
                            if t_span > 1:
                                for off in range(1, t_span): ws.cell(row=f_act + off, column=c_num).value = None
                                try:
                                    ws.merge_cells(start_row=f_act, start_column=c_num, end_row=f_act + t_span - 1, end_column=c_num)
                                    self.excel_reparar_bordes_fusion(ws, f_act, c_num, f_act + t_span - 1, c_num)
                                except: pass

                c_img, t_uni = [], set()
                for r in range(1, ws.max_row + 1):
                    for c in range(1, ws.max_column + 1):
                        celda = ws.cell(row=r, column=c)
                        if isinstance(celda.value, str) and '[' in celda.value:
                            for m in re.findall(r'\[Imagen.*?\]', celda.value, flags=re.IGNORECASE):
                                c_img.append((r, c, m, celda))
                                t_uni.add(m)
                                
                if t_uni:
                    for tag in sorted(list(t_uni)):
                        if tag not in imagenes_cacheadas:
                            ruta_img = filedialog.askopenfilename(title=f"Seleccionar imagen para: {tag} (en {Path(ruta_plantilla).name})", filetypes=[("Imágenes", "*.png *.jpg *.jpeg *.bmp *.gif *.svg")])
                            if ruta_img: 
                                imagenes_cacheadas[tag] = ruta_img
                            
                    for r, c, tag, celda in c_img:
                        ruta = imagenes_cacheadas.get(tag)
                        nv = str(celda.value).replace(tag, "").strip()
                        celda.value = nv if nv else None
                        if ruta:
                            try:
                                try:
                                    img = pxImage(ruta)
                                    img.anchor = celda.coordinate
                                    ws.add_image(img)
                                except Exception as e_img:
                                    if ruta.lower().endswith(".svg"):
                                        from svglib.svglib import svg2rlg
                                        from reportlab.graphics import renderPM
                                        import tempfile
                                        dibujo = svg2rlg(ruta)
                                        temp_png = os.path.join(tempfile.gettempdir(), f"temp_{tag.strip('[]')}.png")
                                        renderPM.drawToFile(dibujo, temp_png, fmt="PNG")
                                        img = pxImage(temp_png)
                                        img.anchor = celda.coordinate
                                        ws.add_image(img)
                                    else: raise e_img
                            except Exception as e:
                                logging.warning(f"Error imagen {ruta}: {e}")

                ruta_salida = Path(ruta_plantilla).parent / f"{Path(ruta_plantilla).stem}_Generado.xlsx"
                wb.save(str(ruta_salida))

            self.after(0, lambda: messagebox.showinfo("Completado", f"¡Se generaron {total_archivos} documento(s) Excel exitosamente!"))

        except Exception as e:
            logging.error(f"Falla crítica en Excel Lote: {e}", exc_info=True)
            self.after(0, lambda: messagebox.showerror("Error", f"Ocurrió un problema:\n{e}"))
        finally:
            self.after(0, lambda: self.btn_excel_procesar.configure(state="normal", text="3. Generar Reportes"))

    # ==========================================
    # MÓDULO 5: CONFIGURACIÓN
    # ==========================================
    def abrir_modulo_config(self):
        for widget in self.frame_main.winfo_children(): widget.destroy()
        ctk.CTkLabel(self.frame_main, text="Configuración del Servidor Ollama (IA)", font=("Segoe UI", 20, "bold")).pack(pady=(20,10), anchor="w", padx=20)
        
        frame_form = ctk.CTkFrame(self.frame_main)
        frame_form.pack(fill="x", padx=20, pady=5)
        
        ctk.CTkLabel(frame_form, text="IP del Servidor:").grid(row=0, column=0, padx=10, pady=10)
        self.ent_ip = ctk.CTkEntry(frame_form)
        self.ent_ip.insert(0, self.ip)
        self.ent_ip.grid(row=0, column=1, padx=10, pady=10)
        
        ctk.CTkLabel(frame_form, text="Puerto:").grid(row=1, column=0, padx=10, pady=10)
        self.ent_puerto = ctk.CTkEntry(frame_form)
        self.ent_puerto.insert(0, self.puerto)
        self.ent_puerto.grid(row=1, column=1, padx=10, pady=10)
        
        ctk.CTkLabel(frame_form, text="Modelo:").grid(row=2, column=0, padx=10, pady=10)
        self.ent_modelo = ctk.CTkEntry(frame_form)
        self.ent_modelo.insert(0, self.modelo)
        self.ent_modelo.grid(row=2, column=1, padx=10, pady=10)

        ctk.CTkButton(self.frame_main, text="Guardar Configuración", command=self.guardar_config).pack(pady=20, padx=20, anchor="w")

    def guardar_config(self):
        self.ip = self.ent_ip.get().strip()
        self.puerto = self.ent_puerto.get().strip()
        self.modelo = self.ent_modelo.get().strip()
        self.cursor.execute("UPDATE configuraciones SET ip=?, puerto=?, modelo=? WHERE id=1", (self.ip, self.puerto, self.modelo))
        self.conn.commit()
        messagebox.showinfo("Éxito", "Configuración guardada correctamente.")

if __name__ == "__main__":
    app = KeyCapitalHub()
    app.mainloop()